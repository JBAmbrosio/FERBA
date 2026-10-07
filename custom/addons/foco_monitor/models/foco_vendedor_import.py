# -*- coding: utf-8 -*-
"""Importador de la Hoja Maestra de un vendedor.

Carga el .xlsx que hoy mantiene cada vendedor a mano y crea/actualiza los
clientes (res.partner) con su geo, la oportunidad (crm.lead) cuando la fila la
trae, y el enlace al expediente. Las coordenadas se SACAN del propio enlace de
Google Maps de la columna de ubicacion: no hay que re-pinchar el mapa.

Mapea las columnas POR EL TEXTO del encabezado (no por su letra), asi que
aguanta que una hoja reordene columnas o cambie de vendedor.
"""

import base64
import io
import re
import unicodedata
from datetime import datetime

from odoo import fields, models
from odoo.exceptions import UserError

# Varias formas en que un enlace de Google Maps trae las coordenadas.
_RE_COORD = [
    re.compile(r'[@/](-?\d{1,3}\.\d{3,}),\s*(-?\d{1,3}\.\d{3,})'),
    re.compile(r'[?&](?:q|ll|daddr|destination|center)='
               r'(-?\d{1,3}\.\d{3,}),\s*(-?\d{1,3}\.\d{3,})'),
    re.compile(r'!3d(-?\d{1,3}\.\d+)!4d(-?\d{1,3}\.\d+)'),
    re.compile(r'(-?\d{1,2}\.\d{4,}),\s*(-?\d{2,3}\.\d{4,})'),
]

_ESTATUS = {
    'SEGUIMIENTO ACTIVO': 'activo',
    'POSTERGADO': 'postergado',
    'VENTA CERRADA': 'cerrado',
}


def _norm(s):
    """Mayusculas, sin acentos y sin espacios de mas (para comparar textos)."""
    if s is None:
        return ''
    s = str(s).strip()
    s = ''.join(c for c in unicodedata.normalize('NFKD', s)
                if not unicodedata.combining(c))
    return re.sub(r'\s+', ' ', s).upper()


def _coords(url):
    """(lat, lon) de un enlace de Google Maps, o None si no las trae."""
    if not url:
        return None
    for rx in _RE_COORD:
        m = rx.search(url)
        if m:
            lat, lon = float(m.group(1)), float(m.group(2))
            if -90 <= lat <= 90 and -180 <= lon <= 180 and (lat or lon):
                return lat, lon
    return None


class FocoVendedorImport(models.TransientModel):
    _name = 'foco.vendedor.import'
    _description = 'Importar Hoja Maestra de vendedor'

    archivo = fields.Binary(string='Hoja Maestra (.xlsx)', required=True,
                            attachment=False)
    nombre_archivo = fields.Char(string='Nombre del archivo')
    vendedor_id = fields.Many2one(
        'res.users', string='Vendedor por defecto',
        default=lambda s: s.env.user,
        help='Se asigna a los clientes cuya columna VENDEDOR no coincida con '
             'un usuario de Odoo.')
    actualizar = fields.Boolean(
        string='Actualizar si ya existe', default=True,
        help='Si un cliente con el mismo nombre ya fue importado, se actualiza '
             'en vez de duplicarlo.')
    crear_oportunidad = fields.Boolean(
        string='Crear oportunidad cuando la fila la traiga', default=True)

    def _col_map(self, headers):
        """De una fila de encabezados -> indice de columna por campo logico."""
        m = {}
        for i, h in enumerate(headers):
            n = _norm(h)
            if not n:
                continue
            if 'CLIENTE' in n or 'RAZON SOCIAL' in n:
                m.setdefault('cliente', i)
            elif 'EMPAQUE' in n:
                m.setdefault('empaque', i)
            elif 'TELEFONO' in n or n == 'TEL':
                m.setdefault('telefono', i)
            elif 'CONTACTO' in n:
                m.setdefault('contacto', i)
            elif 'ESTADO' in n:
                m.setdefault('estado', i)
            elif 'ZONA' in n:
                m.setdefault('zona', i)
            elif 'CULTIVO' in n:
                m.setdefault('cultivo', i)
            elif 'OPORTUNIDAD' in n:
                m.setdefault('oportunidad', i)
            elif 'ESTATUS' in n or 'ESTATU' in n:
                m.setdefault('estatus', i)
            elif 'FECHA' in n:
                m.setdefault('fecha', i)
            elif 'VENDEDOR' in n:
                m.setdefault('vendedor', i)
            elif 'EXPEDIENTE' in n:
                m.setdefault('expediente', i)
            elif 'FISICA' in n:
                m.setdefault('ubicacion_fisica', i)
            elif 'UBICACION' in n or 'UBICAION' in n:
                m.setdefault('ubicacion', i)
        return m

    def action_importar(self):
        self.ensure_one()
        try:
            import openpyxl
        except ImportError:
            raise UserError('Falta la libreria openpyxl en el servidor.')
        if not self.archivo:
            raise UserError('Sube la Hoja Maestra primero.')

        data = base64.b64decode(self.archivo)
        try:
            wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        except Exception as e:
            raise UserError('No pude leer el archivo .xlsx: %s' % e)
        ws = wb[wb.sheetnames[0]]

        # localizar la fila de encabezados (la que trae "CLIENTE")
        cols, hdr_row = None, None
        for r in range(1, min(ws.max_row, 20) + 1):
            fila = [ws.cell(row=r, column=c).value
                    for c in range(1, ws.max_column + 1)]
            mm = self._col_map(fila)
            if 'cliente' in mm:
                cols, hdr_row = mm, r
                break
        if not cols:
            raise UserError(
                'No encontre la columna CLIENTE en las primeras filas del archivo.')

        Partner = self.env['res.partner']
        Users = self.env['res.users']
        # sudo: el administrador de Foco puede no estar en un grupo de Ventas,
        # y aqui creamos la oportunidad con su vendedor explicito.
        Lead = self.env['crm.lead'].sudo()
        State = self.env['res.country.state']
        mx = self.env.ref('base.mx', raise_if_not_found=False)

        cache_user, cache_estado = {}, {}

        def _usuario(nombre):
            n = _norm(nombre)
            if not n:
                return self.vendedor_id
            if n not in cache_user:
                u = Users.search([('name', 'ilike', str(nombre).strip())], limit=1)
                cache_user[n] = u or self.vendedor_id
            return cache_user[n]

        def _estado(nombre):
            if not mx or not nombre:
                return self.env['res.country.state']
            n = _norm(nombre)
            if n not in cache_estado:
                cache_estado[n] = State.search(
                    [('country_id', '=', mx.id),
                     ('name', 'ilike', str(nombre).strip())], limit=1)
            return cache_estado[n]

        def _val(cells, key):
            i = cols.get(key)
            if i is None or i >= len(cells):
                return None
            return cells[i].value

        def _txt(cells, key):
            v = _val(cells, key)
            return str(v).strip() if v not in (None, '') else False

        def _url(cells, key):
            i = cols.get(key)
            if i is None or i >= len(cells):
                return ''
            cell = cells[i]
            if cell.hyperlink and cell.hyperlink.target:
                return cell.hyperlink.target
            if isinstance(cell.value, str) and 'http' in cell.value:
                return cell.value.strip()
            return ''

        creados = actualizados = con_geo = sin_geo = oportunidades = 0

        for r in range(hdr_row + 1, ws.max_row + 1):
            cells = [ws.cell(row=r, column=c)
                     for c in range(1, ws.max_column + 1)]
            nombre = _val(cells, 'cliente')
            tel = _val(cells, 'telefono')
            if not nombre and not tel:
                continue
            nombre = (str(nombre).strip() if nombre else '') or ('Sin nombre %s' % r)

            ub_url = _url(cells, 'ubicacion')
            latlon = _coords(ub_url)
            exp_url = _url(cells, 'expediente')
            estatus = _ESTATUS.get(_norm(_val(cells, 'estatus')), 'prospecto')
            vend = _usuario(_val(cells, 'vendedor'))
            estado = _estado(_val(cells, 'estado'))

            vals = {
                'foco_cliente_ventas': True,
                'foco_empaque': _txt(cells, 'empaque'),
                'foco_contacto': _txt(cells, 'contacto'),
                'foco_zona': _txt(cells, 'zona'),
                'foco_cultivo': _txt(cells, 'cultivo'),
                'foco_estatus': estatus,
                'foco_expediente_url': exp_url or False,
                'foco_ubicacion_fisica': _txt(cells, 'ubicacion_fisica'),
                'foco_ubicacion_url': ub_url or False,
                'phone': _txt(cells, 'telefono'),
                'user_id': vend.id if vend else False,
            }
            if estado:
                vals['state_id'] = estado.id
                if mx:
                    vals['country_id'] = mx.id
            if latlon:
                vals['partner_latitude'] = latlon[0]
                vals['partner_longitude'] = latlon[1]
                vals['foco_geo_pendiente'] = False
                con_geo += 1
            else:
                vals['foco_geo_pendiente'] = bool(ub_url)
                sin_geo += 1

            existente = Partner.search(
                [('foco_cliente_ventas', '=', True), ('name', '=ilike', nombre)],
                limit=1) if self.actualizar else Partner.browse()
            if existente:
                existente.write(vals)
                cliente = existente
                actualizados += 1
            else:
                cliente = Partner.create(
                    dict(vals, name=nombre, company_type='company'))
                creados += 1

            op = _txt(cells, 'oportunidad')
            if op and self.crear_oportunidad:
                fecha = _val(cells, 'fecha')
                deadline = fecha.date() if isinstance(fecha, datetime) else False
                existe = Lead.search(
                    [('partner_id', '=', cliente.id), ('name', '=', op[:200])],
                    limit=1)
                if not existe:
                    Lead.create({
                        'name': op[:200], 'type': 'opportunity',
                        'partner_id': cliente.id,
                        'user_id': vend.id if vend else False,
                        'date_deadline': deadline or False,
                        'description': 'Cultivo: %s' % (_txt(cells, 'cultivo') or '—'),
                    })
                    oportunidades += 1

        msg = ('Importacion lista — %d clientes nuevos, %d actualizados, '
               '%d con ubicacion, %d sin ubicacion, %d oportunidades.'
               % (creados, actualizados, con_geo, sin_geo, oportunidades))
        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': 'Hoja Maestra', 'message': msg,
                'type': 'success', 'sticky': True,
                'next': {'type': 'ir.actions.act_window_close'},
            },
        }
